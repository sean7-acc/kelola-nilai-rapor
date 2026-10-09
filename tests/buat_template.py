"""Membuat berkas contoh template e-rapor untuk pengujian.

Struktur mengikuti Lampiran PRD:
  - judul di baris 1
  - nomor di kolom A, NISN di kolom E, nama di kolom F
  - nilai rapor di kolom G, kolom TP di kolom H dan seterusnya, validasi di kolom M
  - data siswa mulai baris 7
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

TUJUAN = r"C:\xampp\htdocs\kelola nilai rapor\tests\fixture"
os.makedirs(TUJUAN, exist_ok=True)

GARIS = Border(*[Side(style="thin", color="BFBFBF")] * 4)
KEPALA = Font(bold=True, size=10)
ISI = Font(size=10)


def siswa_acak(jumlah, awal_nisn, daftar_nama):
    hasil = []
    for i in range(jumlah):
        nisn = str(awal_nisn + i).zfill(10)
        hasil.append((nisn, daftar_nama[i % len(daftar_nama)] + " " + str(i + 1)))
    return hasil


NAMA = [
    "ADI PRATAMA", "BELLAWATI", "CANDRA WIJAYA", "DINA MARIANA", "EKA SAPUTRA",
    "FITRIANI", "GALIH PERMANA", "HANA LESTARI", "INDRA GUNAWAN", "JULIA SAPUTRI",
    "KURNIAWAN", "LINA HERAWATI", "MAHMUD RIDWAN", "NANDA PRAMESWARI", "OKTA FAJAR",
    "PUTRI ANGGRAINI", "QOMARUDIN", "RATNA DEWI", "SANDI KURNIA", "TANTI SUMARNI",
    "USUP HAMDANI", "VINA OKTAVIA", "WIYONO", "XENA PUSPITA", "YOGA PRADANA",
    "ZAHRA AULIA", "ANDIKA Saputra", "BUNGA MAHARANI", "DIMAS PRAYOGA", "ELSA RISMAWATI",
    "FAJAR NUGROHO", "GITA PERTIWI", "HENDRA KUSUMA", "INTAN PERMATA", "JOKO SUSILO",
    "KIKI AMELIA", "LUKMAN HAKIM", "MEGA WULANDARI", "NICO SANJAYA", "OLIVIA RAHMA",
]


def template_varian_1():
    """Layout sesuai lampiran PRD (judul ganda kata KELAS)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "NILAI"

    judul = "FORMAT IMPORT NILAI RAPOR BAHASA INDONESIA, KELAS KELAS X TSM - REGULER"
    ws.merge_cells("A1:M1")
    ws["A1"] = judul
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "SMK NEGERI 2 KASONGAN"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A3"] = "TAHUN PELAJARAN 2025/2026"
    ws["A3"].font = Font(bold=True, size=11)

    kepala = {
        "A6": "NO", "B6": "KODE", "C6": "ROMBEL", "D6": "JK",
        "E6": "NISN", "F6": "NAMA SISWA", "G6": "NILAI RAPOR",
        "H6": "TP", "I6": "TP", "J6": "TP", "K6": "TP", "L6": "TP",
        "M6": "VALIDASI",
    }
    for sel, teks in kepala.items():
        ws[sel] = teks
        ws[sel].font = KEPALA
        ws[sel].fill = PatternFill("solid", fgColor="DDEBF7")
        ws[sel].border = GARIS
        ws[sel].alignment = Alignment(horizontal="center")

    data = siswa_acak(40, 1012345601, NAMA)
    for i, (nisn, nama) in enumerate(data):
        baris = 7 + i
        ws.cell(baris, 1, i + 1).font = ISI
        ws.cell(baris, 2, "S" + str(i + 1).zfill(3)).font = ISI
        ws.cell(baris, 3, "X TSM").font = ISI
        ws.cell(baris, 4, "L" if i % 2 == 0 else "P").font = ISI
        c_nisn = ws.cell(baris, 5, nisn)
        c_nisn.font = ISI
        c_nisn.number_format = "@"          # teks, agar nol di depan terjaga
        ws.cell(baris, 6, nama).font = ISI
        for kolom in range(7, 14):
            ws.cell(baris, kolom).font = ISI

    for kolom, lebar in zip("ABCDEFGHIJKLM", [6, 9, 10, 6, 13, 30, 13, 7, 7, 7, 7, 7, 12]):
        ws.column_dimensions[kolom].width = lebar

    path = os.path.join(TUJUAN, "template_bahasa_indonesia.xlsx")
    wb.save(path)
    return path, data


def template_varian_2():
    """Layout berbeda: kolom lebih kecil, baris judul kolom berbeda."""
    wb = Workbook()
    ws = wb.active
    ws.title = "IMPORT"

    judul = "FORMAT IMPORT NILAI RAPOR MATEMATIKA, KELAS XI TKJ - REGULER"
    ws.merge_cells("A1:H1")
    ws["A1"] = judul
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")

    kepala = {"A3": "NO", "B3": "NISN", "C3": "NAMA PESERTA DIDIK",
              "D3": "NILAI", "E3": "TP", "F3": "VALIDASI"}
    for sel, teks in kepala.items():
        ws[sel] = teks
        ws[sel].font = KEPALA
        ws[sel].fill = PatternFill("solid", fgColor="E2EFDA")

    data = siswa_acak(35, 1023456701, list(reversed(NAMA)))
    for i, (nisn, nama) in enumerate(data):
        baris = 4 + i
        ws.cell(baris, 1, i + 1)
        c_nisn = ws.cell(baris, 2, nisn)
        c_nisn.number_format = "@"
        ws.cell(baris, 3, nama)

    for kolom, lebar in zip("ABCDEF", [6, 13, 30, 10, 8, 12]):
        ws.column_dimensions[kolom].width = lebar

    path = os.path.join(TUJUAN, "template_matematika.xlsx")
    wb.save(path)
    return path, data


def template_tanpa_kelas():
    """Judul tanpa pemisah dan tanpa kata KELAS — kelas harus diisi manual."""
    wb = Workbook()
    ws = wb.active
    ws.title = "NILAI"

    ws["A1"] = "FORMAT IMPORT NILAI RAPOR PENDIDIKAN AGAMA ISLAM"
    ws["A1"].font = Font(bold=True, size=13)

    ws["A3"] = "NO"
    ws["B3"] = "NISN"
    ws["C3"] = "NAMA"

    data = siswa_acak(10, 1034567801, NAMA)
    for i, (nisn, nama) in enumerate(data):
        baris = 4 + i
        ws.cell(baris, 1, i + 1)
        c_nisn = ws.cell(baris, 2, nisn)
        c_nisn.number_format = "@"
        ws.cell(baris, 3, nama)

    path = os.path.join(TUJUAN, "template_tanpa_kelas.xlsx")
    wb.save(path)
    return path, data


def template_gaya_hubung():
    """Judul dengan pemisah hubung sebelum kata KELAS (pola A2)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "NILAI"

    ws["A1"] = "FORMAT IMPORT NILAI RAPOR BAHASA INGGRIS - KELAS XI AKL - REGULER"
    ws["A1"].font = Font(bold=True, size=13)

    ws["A3"] = "NISN"
    ws["B3"] = "NAMA SISWA"

    data = siswa_acak(12, 1056789001, NAMA)
    for i, (nisn, nama) in enumerate(data):
        baris = 4 + i
        c = ws.cell(baris, 1, nisn)
        c.number_format = "@"
        ws.cell(baris, 2, nama)

    path = os.path.join(TUJUAN, "template_gaya_hubung.xlsx")
    wb.save(path)
    return path, data


def template_duplikat():
    """NISN sama muncul dua kali — harus dilewati dengan peringatan."""
    wb = Workbook()
    ws = wb.active
    ws.title = "NILAI"
    ws["A1"] = "FORMAT IMPORT NILAI RAPOR SEJARAH, KELAS X AKL - REGULER"
    ws["A3"] = "NISN"
    ws["B3"] = "NAMA SISWA"

    data = [
        ("1045678901", "SISWA PERTAMA"),
        ("1045678901", "SISWA PERTAMA GANDA"),
        ("1045678902", "SISWA KEDUA"),
        ("", "TANPA NISN"),
        ("1045678903", "SISWA KETIGA"),
    ]
    for i, (nisn, nama) in enumerate(data):
        baris = 4 + i
        c = ws.cell(baris, 1, nisn)
        c.number_format = "@"
        ws.cell(baris, 2, nama)

    path = os.path.join(TUJUAN, "template_ganda.xlsx")
    wb.save(path)
    return path, data


if __name__ == "__main__":
    for buat in (template_varian_1, template_varian_2, template_tanpa_kelas,
                 template_gaya_hubung, template_duplikat):
        path, data = buat()
        print(os.path.basename(path), "-", len(data), "siswa")
