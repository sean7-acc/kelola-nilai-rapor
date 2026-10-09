"""Uji otomatis Prototype 2 — Kelola Nilai Rapor.

Fokus Proyek 2:
  FR-06 halaman nilai siswa, FR-07 tabel input nilai per kelas,
  FR-08 pengaturan komponen & bobot, FR-09 hitung otomatis,
  FR-10 simpan & bertahan, FR-11 validasi nilai, FR-12 salin clipboard,
  FR-13 ekspor .xlsx, FR-16 impor ulang menjaga nilai.

Menjalankan server statis lokal lalu memeriksa alur dengan Chrome/Edge
(Playwright). Server memakai port 8766 agar dapat berjalan bersamaan
dengan uji Prototype 1.

Pemakaian:  python tests/uji_prototype2.py
"""
import os
import socket
import subprocess
import sys
import tempfile
import time

PROYEK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURE = os.path.join(PROYEK, "tests", "fixture")
PORT = 8766
BASE = "http://127.0.0.1:%d/" % PORT

gagal = []
langkah_ke = [0]


def langkah(teks):
    langkah_ke[0] += 1
    print("\n[%02d] %s" % (langkah_ke[0], teks), flush=True)


def ok(teks):
    print("     OK  " + teks, flush=True)


def periksa(kondisi, teks):
    if kondisi:
        ok(teks)
    else:
        gagal.append(teks)
        print("  GAGAL  " + teks, flush=True)
    return kondisi


def port_bebas(port):
    with socket.socket() as s:
        return s.connect_ex(("127.0.0.1", port)) != 0


def mulai_server():
    proses = subprocess.Popen(
        [sys.executable, "-m", "http.server", str(PORT), "--bind", "127.0.0.1"],
        cwd=PROYEK,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    for _ in range(60):
        if not port_bebas(PORT):
            return proses
        time.sleep(0.2)
    raise RuntimeError("Server statis tidak dapat dijalankan di port %d" % PORT)


def luncurkan(playwright):
    paksa = os.environ.get("BROWSER_KANAL")
    kanal = (paksa,) if paksa else ("chrome", "msedge")
    for nama in kanal:
        try:
            browser = playwright.chromium.launch(channel=nama, headless=True)
            print("     browser: " + nama, flush=True)
            return browser
        except Exception as e:
            if paksa:
                raise
    return playwright.chromium.launch(headless=True)


def km_terbuka(halaman):
    """Baca argumen km pada location.hash yang sedang tampil."""
    return int(halaman.evaluate(
        "(() => new URLSearchParams((location.hash.split('?')[1]||'')).get('km') || 0)()"))


def nilai_masukan(halaman, selektor):
    """Baca nilai seluruh masukan yang cocok dengan selektor."""
    return halaman.locator(selektor).evaluate_all("els => els.map(e => e.value)")


def siapkan_komponen(halaman):
    """Bila halaman pilih belum punya komponen, simpan bawaan 50/25/25."""
    halaman.wait_for_selector("#tombol-simpan, .keadaan-kosong", timeout=8000)
    if halaman.locator(".keadaan-kosong").count():
        halaman.click("a[href^='#/komponen']")
        halaman.wait_for_selector("#simpan-komponen", timeout=8000)
        halaman.click("#simpan-komponen")
    halaman.wait_for_selector(".baris-siswa", timeout=8000)


def impor(halaman, nama_berkas, isi_kelas=""):
    """Impor template lalu konfirmasi dialog. Kembali ke halaman pilih."""
    halaman.goto(BASE + "#/impor")
    halaman.wait_for_selector("#berkas-template", state="attached")
    halaman.set_input_files("#berkas-template", os.path.join(FIXTURE, nama_berkas))
    halaman.wait_for_selector("#area-modal:not([hidden]) #inp-kelas")
    if isi_kelas:
        halaman.fill("#inp-kelas", isi_kelas)
    halaman.click("#area-modal button:has-text('Simpan')")
    halaman.wait_for_selector("#area-modal[hidden]", state="attached", timeout=8000)


def main():
    from playwright.sync_api import sync_playwright

    server = mulai_server()
    kesalahan_konsol = []

    try:
        with sync_playwright() as pw:
            browser = luncurkan(pw)
            konteks = browser.new_context(
                viewport={"width": 1280, "height": 900},
                accept_downloads=True)
            konteks.grant_permissions(["clipboard-read", "clipboard-write"])
            halaman = konteks.new_page()

            def pada_konsol(pesan):
                if pesan.type != "error":
                    return
                lokasi = pesan.location or {}
                url = lokasi.get("url", "") if isinstance(lokasi, dict) else ""
                if "favicon" in url or "favicon" in pesan.text:
                    return
                if ".well-known" in url or ".well-known" in pesan.text:
                    return
                kesalahan_konsol.append("console: " + pesan.text + (" @ " + url if url else ""))

            halaman.on("pageerror", lambda e: kesalahan_konsol.append("pageerror: " + str(e)))
            halaman.on("console", pada_konsol)

            # ---------------------------------------------------------- 1
            langkah("FR-02/FR-07 — impor lalu siapkan komponen, tabel nilai tampil")
            impor(halaman, "template_bahasa_indonesia.xlsx")
            siapkan_komponen(halaman)
            km = km_terbuka(halaman)
            periksa(km == 1, "pasangan pertama memakai km=1 (dapat %s)" % km)
            jumlah = halaman.locator(".baris-siswa").count()
            periksa(jumlah == 40, "40 baris siswa di tabel nilai (dapat %d)" % jumlah)
            periksa(halaman.locator("#tombol-simpan").count() == 1, "tombol Simpan Nilai tampil")
            periksa(halaman.locator("#tombol-salin").count() == 1, "tombol Salin Nilai Rapor tampil")
            periksa(halaman.locator("#tombol-ekspor").count() == 1, "tombol Ekspor .xlsx tampil")
            periksa(halaman.locator(".tabel-nilai th:has-text('HARIAN')").count() == 1,
                    "kolom komponen HARIAN tampil")
            periksa(halaman.locator(".tabel-nilai th:has-text('UTS')").count() == 1,
                    "kolom komponen UTS tampil")
            periksa(halaman.locator(".tabel-nilai th:has-text('UAS')").count() == 1,
                    "kolom komponen UAS tampil")

            # ---------------------------------------------------------- 2
            langkah("FR-08 — komponen bawaan dan indikator jumlah bobot")
            halaman.goto(BASE + "#/komponen?km=%d" % km)
            halaman.wait_for_selector("#simpan-komponen")
            periksa(halaman.locator(".komponen-baris").count() == 3,
                    "3 komponen bawaan terisi")
            isi = nilai_masukan(halaman, ".komponen-baris .inp-kb")
            periksa(isi == ["50", "25", "25"], "bobot bawaan 50/25/25: %r" % isi)
            indikator = halaman.locator("#bobot-ringkas")
            periksa("100%" in indikator.inner_text(), "total bobot 100% tertera")
            periksa("bobot-oke" in indikator.get_attribute("class"),
                    "indikator hijau saat total 100%")

            # Tambah lalu hapus baris tanpa disimpan
            halaman.click("#tambah-komponen")
            baris = halaman.locator(".komponen-baris")
            periksa(baris.count() == 4, "baris komponen baru dapat ditambah")
            baris.last.locator(".inp-kn").fill("PROYEK")
            baris.last.locator(".inp-kb").fill("20")
            periksa("120%" in indikator.inner_text(), "total 120% tercermin")
            baris.last.locator(".hapus-komponen").click()
            periksa(halaman.locator(".komponen-baris").count() == 3,
                    "baris dapat dihapus kembali")

            # Ubah bobot menjadi tidak 100%, simpan (dengan peringatan)
            halaman.locator(".komponen-baris:has(.inp-kn[value='UAS']) .inp-kb").fill("20")
            periksa("belum sama dengan 100%" in indikator.inner_text(),
                    "peringatan saat total bobot bukan 100%")
            halaman.click("#simpan-komponen")
            halaman.wait_for_selector(".baris-siswa", timeout=8000)
            halaman.goto(BASE + "#/komponen?km=%d" % km)
            halaman.wait_for_selector("#simpan-komponen")
            isi = nilai_masukan(halaman, ".komponen-baris .inp-kb")
            periksa("20" in isi, "bobot yang tidak 100% tetap tersimpan: {0}".format(isi))

            # Kembalikan total menjadi 100%
            halaman.locator(".komponen-baris:has(.inp-kn[value='UAS']) .inp-kb").fill("25")
            halaman.click("#simpan-komponen")
            halaman.wait_for_selector(".baris-siswa", timeout=8000)
            periksa(True, "bobot dikembalikan ke 50/25/25")

            # ---------------------------------------------------------- 3
            langkah("FR-07/FR-09 — isi nilai, hitung otomatis")
            kueri = "#/pilih?km=%d" % km
            halaman.goto(BASE + kueri)
            halaman.wait_for_selector(".baris-siswa")
            baris1 = halaman.locator(".baris-siswa").nth(0)
            isi1 = baris1.locator(".inp-nilai")
            isi1.nth(0).fill("100"); isi1.nth(1).fill("80"); isi1.nth(2).fill("80")
            sel1 = baris1.locator(".sel-rapor")
            periksa(sel1.inner_text() == "90", "rapor baris 1 = 90 (Σ(bobot×nilai)÷100)")
            periksa("nilai-oke" in sel1.get_attribute("class"),
                    "rapor lengkap ditandai hijau")

            baris2 = halaman.locator(".baris-siswa").nth(1)
            baris2.locator(".inp-nilai").first.fill("80")
            sel2 = baris2.locator(".sel-rapor")
            periksa(sel2.inner_text() == "—", "komponen belum lengkap = belum dihitung")

            # ---------------------------------------------------------- 4
            langkah("FR-11 — nilai di luar 1–100 ditolak saat simpan")
            isi1.nth(0).fill("150")
            halaman.click("#tombol-simpan")
            halaman.wait_for_selector(
                ".toast:has-text('di luar 1–100')", timeout=5000)
            periksa(True, "peringatan nilai di luar 1–100 tampil")
            isi1.nth(0).fill("100")
            halaman.click("#tombol-simpan")
            halaman.wait_for_selector(".toast:has-text('Nilai tersimpan')", timeout=5000)
            periksa(True, "nilai tersimpan setelah dikoreksi")

            # ---------------------------------------------------------- 5
            langkah("FR-10 — nilai bertahan setelah ditutup dan dibuka")
            halaman.reload(wait_until="load")
            halaman.wait_for_selector(".baris-siswa")
            nilai1 = halaman.locator(".baris-siswa").nth(0).locator(".inp-nilai") \
                .evaluate_all("els => els.map(e => e.value)")
            periksa(nilai1 == ["100", "80", "80"], "nilai baris 1 tersimpan: %r" % nilai1)
            periksa(halaman.locator(".baris-siswa").nth(0).locator(".sel-rapor").inner_text() == "90",
                    "rapor baris 1 tetap 90 setelah muat ulang")
            nilai2 = halaman.locator(".baris-siswa").nth(1).locator(".inp-nilai") \
                .evaluate_all("els => els.map(e => e.value)")
            periksa(nilai2 == ["80", "", ""], "nilai sebagian siswa juga tersimpan: %r" % nilai2)

            # ---------------------------------------------------------- 6
            langkah("FR-09/FR-12 — isi seluruh kelas, salin nilai rapor")
            halaman.evaluate(
                """(() => {
                    document.querySelectorAll('.baris-siswa').forEach((tr, i) => {
                        const v = [80 + (i % 20), 70 + (i % 25), 90 - (i % 10)];
                        const inps = tr.querySelectorAll('.inp-nilai');
                        inps.forEach((inp, j) => {
                            inp.value = v[j];
                            inp.dispatchEvent(new Event('input', { bubbles: true }));
                        });
                    });
                })()"""
            )
            halaman.wait_for_timeout(150)
            harapan = halaman.locator(".baris-siswa .sel-rapor").all_inner_texts()
            periksa(len(harapan) == 40 and all(t.isdigit() for t in harapan),
                    "40 nilai rapor terhitung otomatis")

            halaman.click("#tombol-salin")
            halaman.wait_for_selector(".toast:has-text('disalin ke papan klip')", timeout=5000)
            teks_papan = halaman.evaluate("navigator.clipboard.readText()")
            baris_papan = teks_papan.replace("\r\n", "\n").split("\n")
            periksa(baris_papan == harapan,
                    "isian papan klip sama dengan kolom nilai rapor")
            periksa(len(baris_papan) == 40, "40 baris nilai disalin")

            # ---------------------------------------------------------- 7
            langkah("FR-13 — ekspor .xlsx berisi NISN, nama, dan nilai rapor")
            with halaman.expect_download() as info:
                halaman.click("#tombol-ekspor")
            unduhan = info.value
            periksa(unduhan.suggested_filename.startswith("nilai_rapor_"),
                    "nama berkas ekspor: %s" % unduhan.suggested_filename)
            tujuan = os.path.join(tempfile.gettempdir(), unduhan.suggested_filename)
            unduhan.save_as(tujuan)

            from openpyxl import load_workbook
            wb = load_workbook(tujuan, read_only=True, data_only=True)
            ws = wb["Nilai Rapor"]
            matriks = list(ws.iter_rows(values_only=True))
            periksa(matriks[0] == ("NISN", "Nama Siswa", "Nilai Rapor"),
                    "baris judul ekspor benar: %r" % (matriks[0],))
            periksa(len(matriks) == 41, "40 data siswa + judul: %d baris" % len(matriks))
            periksa(matriks[1][0] == "1012345601", "NISN ikut diekspor")
            cocok = all(str(r[2]) == harapan[i] for i, r in enumerate(matriks[1:]))
            periksa(cocok, "nilai rapor dalam berkas sama dengan aplikasi")

            # ---------------------------------------------------------- 8
            langkah("FR-06 — halaman nilai siswa menyimpan dan menghitung")
            halaman.goto(BASE + kueri)
            halaman.wait_for_selector(".baris-siswa")
            halaman.locator(".baris-siswa").nth(0).locator(".tautan-nama").click()
            halaman.wait_for_selector("h1:has-text('Nilai Siswa')")
            masukan = halaman.locator("input[data-komponen]")
            periksa(masukan.count() == 3, "3 komponen tersedia di halaman siswa")
            periksa(masukan.nth(0).input_value() == "80",
                    "nilai tersimpan tampil kembali di form siswa")
            periksa(masukan.nth(2).input_value() == "90", "nilai UAS tampil")

            masukan.nth(2).fill("85")
            angka = halaman.locator("#rapor-angka")
            periksa(angka.inner_text() == "79", "rapor di halaman siswa diperbarui otomatis (79)")
            halaman.click("#simpan-siswa")
            halaman.wait_for_selector(".toast:has-text('tersimpan')", timeout=5000)

            halaman.goto(BASE + kueri)
            halaman.wait_for_selector(".baris-siswa")
            nilai1 = halaman.locator(".baris-siswa").nth(0).locator(".inp-nilai") \
                .evaluate_all("els => els.map(e => e.value)")
            periksa(nilai1 == ["80", "70", "85"], "ubah dari halaman siswa tersimpan: %r" % nilai1)
            periksa(halaman.locator(".baris-siswa").nth(0).locator(".sel-rapor").inner_text() == "79",
                    "rapor baris 1 mengikuti nilai terbaru")
            wb.close()
            os.remove(tujuan)

            # ---------------------------------------------------------- 9
            langkah("FR-16 — impor ulang tidak menghapus nilai")
            impor(halaman, "template_bahasa_indonesia.xlsx")
            halaman.wait_for_selector(".baris-siswa")
            periksa(halaman.locator(".baris-siswa").count() == 40,
                    "impor ulang tetap 40 siswa")
            nilai1 = halaman.locator(".baris-siswa").nth(0).locator(".inp-nilai") \
                .evaluate_all("els => els.map(e => e.value)")
            periksa(nilai1 == ["80", "70", "85"], "nilai tetap ada setelah impor ulang: %r" % nilai1)

            # ---------------------------------------------------------- 10
            langkah("Tampilan ponsel (390 px) tetap rapi")
            halaman.set_viewport_size({"width": 390, "height": 844})
            for rute, penanda in (("#/pilih?km=%d" % km, "Pilih Kelas"),
                                  ("#/komponen?km=%d" % km, "Atur Komponen & Bobot")):
                halaman.goto(BASE + rute)
                halaman.wait_for_selector("h1:has-text('%s')" % penanda)
                lebar = halaman.evaluate(
                    "[document.documentElement.scrollWidth, window.innerWidth]")
                periksa(lebar[0] <= lebar[1] + 1,
                        "%s tidak meluber horizontal (%d <= %d)" % (rute, lebar[0], lebar[1]))
            halaman.set_viewport_size({"width": 1280, "height": 900})

            # ---------------------------------------------------------- 11
            langkah("Tanpa galat JavaScript di sepanjang uji")
            unik = sorted(set(kesalahan_konsol))
            for k in unik:
                print("     ! " + k, flush=True)
            periksa(len(unik) == 0, "%d galat konsol" % len(unik))

            browser.close()
    finally:
        server.terminate()

    print("\n" + "=" * 60)
    if gagal:
        print("HASIL: %d pemeriksaan GAGAL" % len(gagal))
        for g in gagal:
            print("  - " + g)
        return 1
    print("HASIL: semua pemeriksaan lulus")
    return 0


if __name__ == "__main__":
    sys.exit(main())